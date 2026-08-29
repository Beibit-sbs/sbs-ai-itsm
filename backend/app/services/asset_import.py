from __future__ import annotations

import base64
import io
import json
import uuid
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, inspect, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session

from app.models.asset import Asset
from app.models.asset_import_batch import AssetImportBatch
from app.models.asset_import_row import AssetImportRow
from app.models.ci_class import (
    ConfigurationItemClass,
    ConfigurationItemClassVersion,
)
from app.models.cmdb_reconciliation import (
    CMDBReconciliationRecord,
    CMDBReconciliationRun,
    CMDBSource,
)
from app.models.user import User
from app.services.cmdb_bootstrap import ensure_standard_cmdb_model
from app.services.cmdb_reconciliation import (
    ReconciliationConflict,
    apply_reconciliation,
    preview_reconciliation,
)

FILE_SIZE_LIMIT_BYTES = 5 * 1024 * 1024
DEFAULT_IMPORT_SOURCE = "excel_import"
LIFECYCLE_BY_ASSET_STATUS = {
    "disposed": "DISPOSED",
    "in_stock": "IN_STOCK",
    "in_repair": "MAINTENANCE",
    "maintenance": "MAINTENANCE",
    "inactive": "RETIRED",
    "active": "ACTIVE",
}

TYPE_MAPPING = {
    "МОНИТОР": "monitor",
    "СИСТЕМНЫЙ БЛОК": "desktop",
    "КОМПЬЮТЕР С МОНИТОРОМ": "computer_set",
    "КОМПЛЕКТУЮЩИЕ": "component",
    "МФУ ПРИНТЕР": "mfp_printer",
    "ПРИНТЕР": "printer",
    "СЕТЕВОЕ ОБОРУДОВАНИЕ": "network_equipment",
    "НОУТБУК": "laptop",
    "ПРОЕКТОР": "projector",
    "МОНОБЛОК": "all_in_one",
    "ЭКРАН": "screen",
    "МАКБУК": "laptop",
    "СКАНЕР": "scanner",
    "КУЛЕР": "component",
}


@dataclass(slots=True)
class ParsedExcelRow:
    row_number: int
    raw: dict[str, Any]


class AssetImportService:
    def parse_excel(self, payload: bytes) -> list[ParsedExcelRow]:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            worksheet = workbook["Лист_1"] if "Лист_1" in workbook.sheetnames else workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()

        if not rows:
            return []

        header_idx = self.detect_header_row(rows)
        parsed: list[ParsedExcelRow] = []
        for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            values = list(row)
            if self._is_empty_row(values):
                continue
            parsed.append(
                ParsedExcelRow(
                    row_number=offset,
                    raw={
                        "name": self._read_col(values, 0),
                        "accepted_at": self._read_col(values, 1),
                        "inventory_number": self._read_col(values, 2),
                        "purchase_cost": self._read_col(values, 6),
                        "current_cost": self._read_col(values, 8),
                        "depreciation_amount": self._read_col(values, 9),
                        "residual_value": self._read_col(values, 10),
                        "assigned_to_name": self._read_col(values, 11),
                        "status_m": self._read_col(values, 12),
                        "status_o": self._read_col(values, 14),
                        "original_type": self._read_col(values, 15),
                        "location": self._read_col(values, 17),
                        "purchase_year": self._read_col(values, 18),
                    },
                )
            )
        return parsed

    def detect_header_row(self, rows: list[tuple[Any, ...]]) -> int:
        markers = ("НАИМЕНОВАНИЕ", "ИНВЕНТАР", "ПЕРВОНАЧАЛЬНАЯ", "КАБИНЕТ", "ТЕХНИКА")
        for idx, row in enumerate(rows[:30]):
            joined = " ".join(self._normalize_text(value) for value in row if value is not None)
            score = sum(1 for marker in markers if marker in joined)
            if score >= 3:
                return idx
        return 6

    def normalize_row(self, raw: dict[str, Any]) -> dict[str, Any]:
        inventory_number = self._clean_str(raw.get("inventory_number"))
        name = self._clean_str(raw.get("name")) or "Imported asset"
        original_type = self._clean_str(raw.get("original_type"))
        location = self._clean_str(raw.get("location"))
        assigned_to_name = self._clean_str(raw.get("assigned_to_name"))

        status_blob = " ".join(
            [
                self._normalize_text(raw.get("status_m")),
                self._normalize_text(raw.get("status_o")),
            ]
        ).strip()

        if "СПИСАНО" in status_blob:
            status = "disposed"
        elif "ПЕРЕВЕДЕНО" in status_blob and "ЗАПАС" in status_blob:
            status = "in_stock"
        else:
            status = "active"

        verification_status = "verified"
        if not location:
            verification_status = "needs_location"

        accepted_at = self._parse_datetime(raw.get("accepted_at"))
        purchase_year = self._parse_year(raw.get("purchase_year"))
        if purchase_year is None and accepted_at is not None:
            purchase_year = accepted_at.year

        normalized = {
            "name": name,
            "inventory_number": inventory_number,
            "asset_type": self._normalize_asset_type(original_type),
            "original_type": original_type,
            "status": status,
            "assigned_to_name": assigned_to_name,
            "location": location or "Location unknown",
            "purchase_year": purchase_year,
            "accepted_at": accepted_at.isoformat() if accepted_at else None,
            "purchase_cost": self._parse_float(raw.get("purchase_cost")),
            "current_cost": self._parse_float(raw.get("current_cost")),
            "depreciation_amount": self._parse_float(raw.get("depreciation_amount")),
            "residual_value": self._parse_float(raw.get("residual_value")),
            "verification_status": verification_status,
            "source": DEFAULT_IMPORT_SOURCE,
        }
        return normalized

    def validate_row(self, normalized: dict[str, Any]) -> str | None:
        if not normalized.get("inventory_number"):
            return "missing_inventory_number"
        return None

    def preview_import(self, db: Session, *, batch: AssetImportBatch, dry_run: bool) -> dict[str, Any]:
        payload = self._extract_payload(batch)
        parsed_rows = self.parse_excel(payload)
        db.execute(delete(AssetImportRow).where(AssetImportRow.batch_id == batch.id))

        existing_inventory = {
            item.inventory_number
            for item in db.scalars(
                select(Asset).where(Asset.tenant_id == batch.tenant_id, Asset.inventory_number.is_not(None))
            ).all()
            if item.inventory_number
        }
        seen_inventory: set[str] = set()

        for entry in parsed_rows:
            normalized = self.normalize_row(entry.raw)
            error_message = self.validate_row(normalized)
            row_status = "valid"
            if error_message:
                row_status = "error"
            else:
                inventory_number = str(normalized["inventory_number"])
                if inventory_number in seen_inventory:
                    row_status = "duplicate"
                    error_message = "duplicate_inventory_number_in_file"
                elif inventory_number in existing_inventory:
                    row_status = "update_candidate"
                seen_inventory.add(inventory_number)

            db.add(
                AssetImportRow(
                    id=self._uuid(),
                    tenant_id=batch.tenant_id,
                    batch_id=batch.id,
                    row_number=entry.row_number,
                    raw_json=json.dumps(entry.raw, ensure_ascii=False),
                    normalized_json=json.dumps(normalized, ensure_ascii=False),
                    status=row_status,
                    error_message=error_message,
                )
            )

        db.flush()
        rows = db.scalars(select(AssetImportRow).where(AssetImportRow.batch_id == batch.id).order_by(AssetImportRow.row_number.asc())).all()
        source, actor = self._reconciliation_source(db, batch)
        reconciliation_records = [
            self._to_reconciliation_record(
                self._parse_json(row.normalized_json)
            )
            for row in rows
            if row.status not in {"error", "duplicate"}
        ]
        runs: list[CMDBReconciliationRun] = []
        outcomes: dict[str, CMDBReconciliationRecord] = {}
        for chunk_index, offset in enumerate(
            range(0, len(reconciliation_records), 500),
            start=1,
        ):
            run = preview_reconciliation(
                db,
                source=source,
                idempotency_key=f"excel-import:{batch.id}:{chunk_index}",
                records=reconciliation_records[offset : offset + 500],
                actor_id=actor.id,
            )
            runs.append(run)
            outcomes.update(
                {
                    record.external_id: record
                    for record in db.scalars(
                        select(CMDBReconciliationRecord).where(
                            CMDBReconciliationRecord.run_id == run.id
                        )
                    ).all()
                }
            )
        for row in rows:
            if row.status in {"error", "duplicate"}:
                continue
            normalized = self._parse_json(row.normalized_json)
            external_id = str(normalized.get("inventory_number", ""))
            record = outcomes.get(external_id)
            if record is None:
                row.status = "error"
                row.error_message = "cmdb_reconciliation_record_missing"
                continue
            if record.outcome == "CREATE":
                row.status = "valid"
                row.error_message = None
            elif record.outcome in {
                "UPDATE",
                "UNCHANGED",
                "SKIPPED",
                "APPLIED_CREATED",
                "APPLIED_UPDATED",
            }:
                row.status = "update_candidate"
                details = self._parse_json(record.normalized_json)
                blocked = details.get("_blocked_fields") or []
                row.error_message = (
                    f"protected_fields:{','.join(str(item) for item in blocked)}"
                    if blocked
                    else None
                )
            else:
                row.status = "error"
                errors = self._parse_list(record.errors_json)
                candidates = self._parse_list(record.candidate_ids_json)
                messages = [str(item) for item in errors]
                if candidates:
                    messages.append(
                        "ambiguous_candidates:"
                        + ",".join(str(item) for item in candidates)
                    )
                row.error_message = ";".join(messages) or (
                    f"cmdb_reconciliation_{record.outcome.lower()}"
                )
        summary = self.build_summary(rows)
        summary["dry_run"] = dry_run
        summary["batch_id"] = batch.id
        run_ids = [run.id for run in runs]
        summary["cmdb_reconciliation_run_id"] = run_ids[0] if run_ids else None
        summary["cmdb_reconciliation_run_ids"] = run_ids
        summary["cmdb_reconciliation_status"] = (
            "PREVIEWED" if runs else "NO_VALID_RECORDS"
        )
        summary["cmdb_ambiguous_rows"] = sum(
            run.ambiguous_count for run in runs
        )
        summary["cmdb_invalid_rows"] = sum(run.invalid_count for run in runs)

        metadata = self._parse_summary(batch.summary_json)
        metadata["preview"] = summary
        metadata["cmdb_reconciliation_run_id"] = run_ids[0] if run_ids else None
        metadata["cmdb_reconciliation_run_ids"] = run_ids
        batch.summary_json = json.dumps(metadata, ensure_ascii=False)
        batch.status = "preview_ready"
        batch.total_rows = summary["total_rows"]
        batch.valid_rows = summary["valid_rows"]
        batch.error_rows = summary["error_rows"]
        batch.skipped_rows = summary["duplicate_rows"]
        db.flush()
        return summary

    def commit_import(self, db: Session, *, batch: AssetImportBatch, dry_run: bool) -> dict[str, Any]:
        rows = db.scalars(select(AssetImportRow).where(AssetImportRow.batch_id == batch.id).order_by(AssetImportRow.row_number.asc())).all()
        imported_rows = 0
        skipped_rows = sum(row.status == "duplicate" for row in rows)
        metadata = self._parse_summary(batch.summary_json)
        raw_run_ids = metadata.get("cmdb_reconciliation_run_ids")
        if isinstance(raw_run_ids, list):
            run_ids = [str(item) for item in raw_run_ids if str(item)]
        else:
            fallback_id = str(
                metadata.get("cmdb_reconciliation_run_id") or ""
            )
            run_ids = [fallback_id] if fallback_id else []
        runs = [
            run
            for run_id in run_ids
            if (run := db.get(CMDBReconciliationRun, run_id)) is not None
        ]
        if (
            not runs
            or len(runs) != len(run_ids)
            or any(run.tenant_id != batch.tenant_id for run in runs)
        ):
            raise ValueError("asset_import_reconciliation_preview_required")
        if any(run.invalid_count or run.ambiguous_count for run in runs):
            raise ReconciliationConflict(
                "asset import contains invalid or ambiguous CMDB records"
            )

        if dry_run:
            imported_rows = sum(
                row.status in {"valid", "update_candidate"}
                for row in rows
            )
        else:
            _, actor = self._reconciliation_source(db, batch)
            applied_runs = [
                apply_reconciliation(
                    db,
                    run_id=run.id,
                    actor_id=actor.id,
                )
                for run in runs
            ]
            outcomes: dict[str, CMDBReconciliationRecord] = {}
            for run in applied_runs:
                outcomes.update(
                    {
                        record.external_id: record
                        for record in db.scalars(
                            select(CMDBReconciliationRecord).where(
                                CMDBReconciliationRecord.run_id == run.id
                            )
                        ).all()
                    }
                )
            runs = applied_runs
            for row in rows:
                if row.status in {"error", "duplicate"}:
                    continue
                normalized = self._parse_json(row.normalized_json)
                external_id = str(normalized.get("inventory_number", ""))
                record = outcomes.get(external_id)
                if record is None:
                    row.status = "error"
                    row.error_message = "cmdb_reconciliation_record_missing"
                    continue
                row.asset_id = record.matched_ci_id
                row.error_message = None
                if record.outcome == "APPLIED_CREATED":
                    row.status = "imported"
                    imported_rows += 1
                elif record.outcome == "APPLIED_UPDATED":
                    row.status = "updated"
                    imported_rows += 1
                elif record.outcome == "UNCHANGED":
                    row.status = "unchanged"
                    skipped_rows += 1
                elif record.outcome == "SKIPPED":
                    row.status = "protected"
                    skipped_rows += 1
                else:
                    row.status = "error"
                    row.error_message = (
                        f"cmdb_reconciliation_{record.outcome.lower()}"
                    )

        summary = self.build_summary(rows)
        summary["dry_run"] = dry_run
        summary["imported_rows"] = imported_rows
        summary["skipped_rows"] = skipped_rows
        summary["error_rows"] = summary["error_rows"]
        summary["cmdb_reconciliation_run_id"] = runs[0].id
        summary["cmdb_reconciliation_run_ids"] = [run.id for run in runs]
        summary["cmdb_reconciliation_status"] = (
            "COMPLETED"
            if all(run.status == "COMPLETED" for run in runs)
            else "COMPLETED_WITH_ERRORS"
        )

        if not dry_run:
            batch.status = "committed"
            batch.imported_rows = imported_rows
            batch.skipped_rows = summary["skipped_rows"]
            batch.error_rows = summary["error_rows"]
            batch.completed_at = datetime.now(UTC)
            metadata["commit"] = summary
            batch.summary_json = json.dumps(metadata, ensure_ascii=False)
            db.flush()
        else:
            batch.status = "preview_ready"
            db.rollback()
        return summary

    def _reconciliation_source(
        self,
        db: Session,
        batch: AssetImportBatch,
    ) -> tuple[CMDBSource, User]:
        if not batch.tenant_id:
            raise ValueError("asset_import_tenant_required")
        source = db.scalar(
            select(CMDBSource).where(
                CMDBSource.tenant_id == batch.tenant_id,
                CMDBSource.code == "EXCEL_ASSET_IMPORT",
            )
        )
        if source is None:
            ensure_standard_cmdb_model(db, batch.tenant_id)
            source = db.scalar(
                select(CMDBSource).where(
                    CMDBSource.tenant_id == batch.tenant_id,
                    CMDBSource.code == "EXCEL_ASSET_IMPORT",
                )
            )
        if source is None:
            raise ValueError("asset_import_cmdb_source_unavailable")
        actor = (
            db.scalar(
                select(User).where(
                    User.tenant_id == batch.tenant_id,
                    User.email == batch.created_by,
                    User.is_active.is_(True),
                )
            )
            if batch.created_by
            else None
        )
        if actor is None:
            actor = db.scalar(
                select(User)
                .where(
                    User.tenant_id == batch.tenant_id,
                    User.is_active.is_(True),
                    User.is_root.is_(False),
                )
                .order_by(User.created_at, User.id)
            )
        if actor is None:
            raise ValueError("asset_import_actor_unavailable")
        return source, actor

    def _to_reconciliation_record(
        self,
        normalized: dict[str, Any],
    ) -> dict[str, Any]:
        inventory_number = self._clean_str(normalized.get("inventory_number"))
        if not inventory_number:
            raise ValueError("asset_import_inventory_number_required")
        return {
            "external_id": inventory_number,
            "name": self._clean_str(normalized.get("name")) or "Imported asset",
            "inventory_number": inventory_number,
            "original_type": self._clean_str(normalized.get("original_type")),
            "lifecycle_status": LIFECYCLE_BY_ASSET_STATUS.get(
                self._clean_str(normalized.get("status")) or "active",
                "ACTIVE",
            ),
            "location": self._clean_str(normalized.get("location"))
            or "Location unknown",
            "assigned_to_name": self._clean_str(
                normalized.get("assigned_to_name")
            ),
            "purchase_date": normalized.get("accepted_at"),
            "purchase_cost": self._parse_float(
                normalized.get("purchase_cost")
            ),
            "current_cost": self._parse_float(
                normalized.get("current_cost")
            ),
            "depreciation_amount": self._parse_float(
                normalized.get("depreciation_amount")
            ),
            "residual_value": self._parse_float(
                normalized.get("residual_value")
            ),
            "purchase_year": self._parse_year(
                normalized.get("purchase_year")
            ),
            "verification_status": self._clean_str(
                normalized.get("verification_status")
            ),
            "attributes": {},
        }

    def upsert_asset(self, db: Session, *, batch: AssetImportBatch, normalized: dict[str, Any]) -> tuple[Asset, str]:
        inventory_number = str(normalized["inventory_number"])
        existing = db.scalar(
            select(Asset).where(Asset.tenant_id == batch.tenant_id, Asset.inventory_number == inventory_number)
        )

        accepted_at = self._parse_datetime(normalized.get("accepted_at"))
        assigned_to_name = self._clean_str(normalized.get("assigned_to_name"))
        location = self._clean_str(normalized.get("location")) or "Location unknown"
        status = self._clean_str(normalized.get("status")) or "active"
        generic_class, generic_version = self._generic_class_snapshot(
            db,
            batch.tenant_id,
        )

        if existing is None:
            asset = Asset(
                id=self._uuid(),
                tenant_id=batch.tenant_id,
                ci_class_id=generic_class.id,
                ci_class_version_id=generic_version.id,
                ci_class_code=generic_class.code,
                ci_class_name=generic_class.name,
                ci_schema_version=generic_version.version,
                ci_schema_hash=generic_version.schema_hash,
                ci_attributes_json="{}",
                lifecycle_status=LIFECYCLE_BY_ASSET_STATUS.get(
                    status,
                    "ACTIVE",
                ),
                owner_user_id=None,
                support_group="Service Desk",
                criticality="MEDIUM",
                environment="OTHER",
                ci_version=1,
                asset_tag=f"IMP-{self._uuid()[:8].upper()}",
                name=self._clean_str(normalized.get("name")) or "Imported asset",
                asset_type=self._clean_str(normalized.get("asset_type")) or "other",
                type=self._clean_str(normalized.get("asset_type")) or "other",
                original_type=self._clean_str(normalized.get("original_type")),
                serial_number=None,
                inventory_number=inventory_number,
                source=DEFAULT_IMPORT_SOURCE,
                source_batch_id=batch.id,
                manufacturer=None,
                model=None,
                status=status,
                owner_name=assigned_to_name or "Warehouse",
                assigned_to_name=assigned_to_name,
                assigned_to_email=None,
                department=None,
                location=location,
                purchase_date=accepted_at,
                accepted_at=accepted_at,
                warranty_until=None,
                purchase_cost=self._parse_float(normalized.get("purchase_cost")),
                current_cost=self._parse_float(normalized.get("current_cost")),
                depreciation_amount=self._parse_float(normalized.get("depreciation_amount")),
                residual_value=self._parse_float(normalized.get("residual_value")),
                purchase_year=self._parse_year(normalized.get("purchase_year")),
                verification_status=self._clean_str(normalized.get("verification_status")) or "verified",
                imported_at=datetime.now(UTC),
                condition="good",
                description="Imported from Excel",
            )
            db.add(asset)
            db.flush()
            return asset, "created"

        governed_class = bool(
            existing.ci_class_id
            and existing.ci_class_code
            and existing.ci_class_code != "GENERIC_ASSET"
        )
        if not existing.ci_class_id:
            existing.ci_class_id = generic_class.id
            existing.ci_class_version_id = generic_version.id
            existing.ci_class_code = generic_class.code
            existing.ci_class_name = generic_class.name
            existing.ci_schema_version = generic_version.version
            existing.ci_schema_hash = generic_version.schema_hash
            existing.ci_attributes_json = "{}"
        existing.name = self._clean_str(normalized.get("name")) or existing.name
        if not governed_class:
            existing.asset_type = (
                self._clean_str(normalized.get("asset_type"))
                or existing.asset_type
            )
            existing.type = (
                self._clean_str(normalized.get("asset_type"))
                or existing.type
            )
        existing.original_type = self._clean_str(normalized.get("original_type")) or existing.original_type
        existing.status = status
        existing.lifecycle_status = LIFECYCLE_BY_ASSET_STATUS.get(
            status,
            existing.lifecycle_status,
        )
        existing.assigned_to_name = assigned_to_name or existing.assigned_to_name
        existing.owner_name = assigned_to_name or existing.owner_name
        existing.location = location
        existing.purchase_date = accepted_at or existing.purchase_date
        existing.accepted_at = accepted_at or existing.accepted_at
        existing.purchase_cost = self._parse_float(normalized.get("purchase_cost"))
        existing.current_cost = self._parse_float(normalized.get("current_cost"))
        existing.depreciation_amount = self._parse_float(normalized.get("depreciation_amount"))
        existing.residual_value = self._parse_float(normalized.get("residual_value"))
        existing.purchase_year = self._parse_year(normalized.get("purchase_year"))
        existing.verification_status = self._clean_str(normalized.get("verification_status")) or existing.verification_status
        existing.source = DEFAULT_IMPORT_SOURCE
        existing.source_batch_id = batch.id
        existing.imported_at = datetime.now(UTC)
        existing.ci_version += 1
        existing.updated_at = datetime.now(UTC)
        db.flush()
        return existing, "updated"

    def _generic_class_snapshot(
        self,
        db: Session,
        tenant_id: str | None,
    ) -> tuple[ConfigurationItemClass, ConfigurationItemClassVersion]:
        if not tenant_id:
            raise ValueError("asset_import_tenant_required")
        cache_key = f"cmdb_generic_snapshot:{tenant_id}"
        cached = db.info.get(cache_key)
        if (
            isinstance(cached, tuple)
            and len(cached) == 2
            and isinstance(cached[0], ConfigurationItemClass)
            and isinstance(cached[1], ConfigurationItemClassVersion)
        ):
            return cached
        generic_class = db.scalar(
            select(ConfigurationItemClass).where(
                ConfigurationItemClass.tenant_id == tenant_id,
                ConfigurationItemClass.code == "GENERIC_ASSET",
            )
        )
        if generic_class is None:
            ensure_standard_cmdb_model(db, tenant_id)
            generic_class = db.scalar(
                select(ConfigurationItemClass).where(
                    ConfigurationItemClass.tenant_id == tenant_id,
                    ConfigurationItemClass.code == "GENERIC_ASSET",
                )
            )
        if generic_class is None:
            raise ValueError("asset_import_generic_ci_class_unavailable")
        version = db.scalar(
            select(ConfigurationItemClassVersion)
            .where(
                ConfigurationItemClassVersion.ci_class_id == generic_class.id,
                ConfigurationItemClassVersion.status == "PUBLISHED",
            )
            .order_by(ConfigurationItemClassVersion.version.desc())
        )
        if version is None:
            raise ValueError("asset_import_generic_ci_schema_unavailable")
        snapshot = (generic_class, version)
        db.info[cache_key] = snapshot
        return snapshot

    def build_summary(self, rows: list[AssetImportRow]) -> dict[str, Any]:
        normalized_rows = [self._parse_json(row.normalized_json) for row in rows]
        return {
            "total_rows": len(rows),
            "valid_rows": sum(
                1
                for row in rows
                if row.status in {
                    "valid",
                    "update_candidate",
                    "imported",
                    "updated",
                    "unchanged",
                    "protected",
                }
            ),
            "error_rows": sum(1 for row in rows if row.status == "error"),
            "duplicate_rows": sum(1 for row in rows if row.status == "duplicate"),
            "update_rows": sum(
                1
                for row in rows
                if row.status in {
                    "update_candidate",
                    "updated",
                    "unchanged",
                    "protected",
                }
            ),
            "disposed_rows": sum(1 for item in normalized_rows if isinstance(item, dict) and item.get("status") == "disposed"),
            "in_stock_rows": sum(1 for item in normalized_rows if isinstance(item, dict) and item.get("status") == "in_stock"),
            "missing_location_rows": sum(
                1 for item in normalized_rows if isinstance(item, dict) and item.get("verification_status") == "needs_location"
            ),
        }

    def _extract_payload(self, batch: AssetImportBatch) -> bytes:
        metadata = self._parse_summary(batch.summary_json)
        encoded = metadata.get("uploaded_file_b64")
        if not isinstance(encoded, str) or not encoded:
            raise ValueError("batch_payload_not_found")
        return base64.b64decode(encoded)

    def _parse_summary(self, raw: str | None) -> dict[str, Any]:
        if not raw:
            return {}
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError:
            return {}
        return parsed if isinstance(parsed, dict) else {}

    def _normalize_asset_type(self, value: str | None) -> str:
        normalized = self._normalize_text(value)
        if not normalized:
            return "other"
        for key, mapped in TYPE_MAPPING.items():
            if key in normalized:
                return mapped
        return "other"

    def _is_empty_row(self, values: list[Any]) -> bool:
        return all(self._clean_str(item) is None for item in values)

    def _read_col(self, values: list[Any], idx: int) -> Any:
        if idx >= len(values):
            return None
        return values[idx]

    def _normalize_text(self, value: Any) -> str:
        text = self._clean_str(value)
        if not text:
            return ""
        return text.upper()

    def _clean_str(self, value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text if text else None

    def _parse_float(self, value: Any) -> float | None:
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, Decimal):
            return float(value)
        text = str(value).strip().replace(" ", "").replace(",", ".")
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            return None

    def _parse_year(self, value: Any) -> int | None:
        if value is None:
            return None
        if isinstance(value, int):
            return value
        text = self._clean_str(value)
        if not text:
            return None
        digits = "".join(ch for ch in text if ch.isdigit())
        if len(digits) == 4:
            return int(digits)
        return None

    def _parse_datetime(self, value: Any) -> datetime | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            if value.tzinfo is None:
                return value.replace(tzinfo=UTC)
            return value.astimezone(UTC)
        if isinstance(value, date):
            return datetime(value.year, value.month, value.day, tzinfo=UTC)
        text = self._clean_str(value)
        if not text:
            return None
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(text, fmt).replace(tzinfo=UTC)
            except ValueError:
                continue
        return None

    def _parse_json(self, raw: str | None) -> dict[str, Any]:
        if not raw:
            return {}
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError:
            return {}
        return parsed if isinstance(parsed, dict) else {}

    def _parse_list(self, raw: str | None) -> list[Any]:
        if not raw:
            return []
        try:
            parsed = json.loads(raw)
        except json.JSONDecodeError:
            return []
        return parsed if isinstance(parsed, list) else []

    def _uuid(self) -> str:
        return str(uuid.uuid4())


def ensure_asset_import_schema(engine: Engine) -> None:
    inspector = inspect(engine)
    if not inspector.has_table("assets"):
        return

    with engine.begin() as connection:
        existing_assets = {column["name"] for column in inspector.get_columns("assets")}
        asset_columns = {
            "source": "ALTER TABLE assets ADD COLUMN source VARCHAR(64)",
            "source_batch_id": "ALTER TABLE assets ADD COLUMN source_batch_id VARCHAR(36)",
            "original_type": "ALTER TABLE assets ADD COLUMN original_type VARCHAR(120)",
            "purchase_cost": "ALTER TABLE assets ADD COLUMN purchase_cost FLOAT",
            "current_cost": "ALTER TABLE assets ADD COLUMN current_cost FLOAT",
            "depreciation_amount": "ALTER TABLE assets ADD COLUMN depreciation_amount FLOAT",
            "residual_value": "ALTER TABLE assets ADD COLUMN residual_value FLOAT",
            "purchase_year": "ALTER TABLE assets ADD COLUMN purchase_year INTEGER",
            "accepted_at": "ALTER TABLE assets ADD COLUMN accepted_at TIMESTAMP",
            "verification_status": "ALTER TABLE assets ADD COLUMN verification_status VARCHAR(64)",
            "imported_at": "ALTER TABLE assets ADD COLUMN imported_at TIMESTAMP",
            "building": "ALTER TABLE assets ADD COLUMN building VARCHAR(120)",
            "floor": "ALTER TABLE assets ADD COLUMN floor VARCHAR(64)",
            "room": "ALTER TABLE assets ADD COLUMN room VARCHAR(120)",
            "location_label": "ALTER TABLE assets ADD COLUMN location_label VARCHAR(255)",
            "location_verified_at": "ALTER TABLE assets ADD COLUMN location_verified_at TIMESTAMP",
            "responsible_person_name": "ALTER TABLE assets ADD COLUMN responsible_person_name VARCHAR(200)",
            "responsible_person_position": "ALTER TABLE assets ADD COLUMN responsible_person_position VARCHAR(200)",
            "responsible_department": "ALTER TABLE assets ADD COLUMN responsible_department VARCHAR(200)",
            "mol_name": "ALTER TABLE assets ADD COLUMN mol_name VARCHAR(200)",
            "mol_department": "ALTER TABLE assets ADD COLUMN mol_department VARCHAR(200)",
            "initial_cost": "ALTER TABLE assets ADD COLUMN initial_cost FLOAT",
            "residual_cost": "ALTER TABLE assets ADD COLUMN residual_cost FLOAT",
            "writeoff_date": "ALTER TABLE assets ADD COLUMN writeoff_date TIMESTAMP",
            "writeoff_reason": "ALTER TABLE assets ADD COLUMN writeoff_reason TEXT",
            "assigned_at": "ALTER TABLE assets ADD COLUMN assigned_at TIMESTAMP",
            "moved_at": "ALTER TABLE assets ADD COLUMN moved_at TIMESTAMP",
            "disposed_at": "ALTER TABLE assets ADD COLUMN disposed_at TIMESTAMP",
            "last_inventory_at": "ALTER TABLE assets ADD COLUMN last_inventory_at TIMESTAMP",
            "last_verified_at": "ALTER TABLE assets ADD COLUMN last_verified_at TIMESTAMP",
            "notes": "ALTER TABLE assets ADD COLUMN notes TEXT",
        }
        for column_name, ddl in asset_columns.items():
            if column_name not in existing_assets:
                connection.exec_driver_sql(ddl)

        if not inspector.has_table("asset_history"):
            connection.exec_driver_sql(
                """
                CREATE TABLE asset_history (
                    id VARCHAR(36) PRIMARY KEY,
                    asset_id VARCHAR(36) NOT NULL,
                    actor_id VARCHAR(36),
                    action VARCHAR(80) NOT NULL,
                    old_value TEXT,
                    new_value TEXT,
                    comment TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP NOT NULL,
                    FOREIGN KEY(asset_id) REFERENCES assets(id) ON DELETE CASCADE,
                    FOREIGN KEY(actor_id) REFERENCES users(id) ON DELETE SET NULL
                )
                """
            )
