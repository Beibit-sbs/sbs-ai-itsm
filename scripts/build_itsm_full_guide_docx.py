from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "manuals"
OUTPUT = OUTPUT_DIR / "SBS-AI-ITSM-FULL-GUIDE-RU-v0.1.0.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
MUTED = "5F6B76"
GREEN = "1F6B4F"
GOLD = "7A5A00"
RED = "9B1C1C"
INK = "1F2933"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table, color="C8D0D9", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def create_numbering(document: Document, kind: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def configure_document(document: Document) -> tuple[int, int]:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("SBS AI ITSM  |  Полное руководство")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    run = footer_p.add_run("Версия 0.1.0  |  Стр. ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(footer_p, "PAGE")

    first_footer = section.first_page_footer
    first_p = first_footer.paragraphs[0]
    first_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = first_p.add_run("SBS AI ITSM · Smart Business Systems · 2026")
    set_run_font(run, size=8.5, color=MUTED)

    bullet_num_id = create_numbering(document, "bullet")
    decimal_num_id = create_numbering(document, "decimal")
    return bullet_num_id, decimal_num_id


def add_paragraph(document, text="", *, bold_prefix=None, italic=False, color=None, align=None, after=6):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=color)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, italic=italic, color=color)
    else:
        run = p.add_run(text)
        set_run_font(run, italic=italic, color=color)
    return p


def add_list_item(document, text: str, num_id: int, *, bold_prefix=None):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        set_run_font(p.add_run(text))
    return p


def add_bullets(document, items, bullet_num_id):
    for item in items:
        if isinstance(item, tuple):
            add_list_item(document, item[0] + item[1], bullet_num_id, bold_prefix=item[0])
        else:
            add_list_item(document, item, bullet_num_id)


def add_steps(document, items, decimal_num_id):
    for item in items:
        add_list_item(document, item, decimal_num_id)


def add_heading(document, text, level=1, *, new_page=False):
    if new_page:
        document.add_page_break()
    p = document.add_paragraph(text, style=f"Heading {level}")
    return p


def add_callout(document, title: str, text: str, *, kind="info"):
    color = {"info": BLUE, "success": GREEN, "warning": GOLD, "danger": RED}.get(kind, BLUE)
    fill = {"info": CALLOUT, "success": "EAF5EF", "warning": "FFF7E0", "danger": "FBEAEC"}.get(kind, CALLOUT)
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=color, size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    lead = p.add_run(title + ": ")
    set_run_font(lead, bold=True, color=color)
    body = p.add_run(text)
    set_run_font(body, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(document, lines: list[str]):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color="D5D9DE", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=9, color=INK)
        if index < len(lines) - 1:
            run.add_break()
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(document, headers, rows, widths_dxa, *, font_size=9.2):
    table = document.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(str(value))
        set_run_font(run, size=font_size, color=NAVY, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=INK)
    set_table_geometry(table, widths_dxa)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_label_value(document, label: str, value: str):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    lead = p.add_run(label + ": ")
    set_run_font(lead, bold=True, color=DARK_BLUE)
    set_run_font(p.add_run(value))


def add_cover(document):
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(66)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    set_run_font(kicker.add_run("SMART BUSINESS SYSTEMS"), size=10.5, color=BLUE, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("SBS AI ITSM"), size=32, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    set_run_font(subtitle.add_run("Полное руководство пользователя, администратора и внедрения"), size=16, color=DARK_BLUE)

    desc = document.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.paragraph_format.space_after = Pt(44)
    set_run_font(desc.add_run("Назначение платформы · ежедневная работа · администрирование · AI · перенос на сервер"), size=10.5, color=MUTED, italic=True)

    table = document.add_table(rows=4, cols=2)
    set_table_geometry(table, [2700, 6660])
    set_table_borders(table, color="D5DDE5")
    for row, (label, value) in zip(table.rows, [
        ("Документ", "Единое эксплуатационное руководство"),
        ("Версия системы", "0.1.0"),
        ("Контур", "Локальная Windows-среда; готово для демонстрации и UAT"),
        ("Дата редакции", "13 августа 2026 года"),
    ]):
        set_cell_shading(row.cells[0], PALE_BLUE)
        p0 = row.cells[0].paragraphs[0]
        set_run_font(p0.add_run(label), size=9.5, color=NAVY, bold=True)
        p1 = row.cells[1].paragraphs[0]
        set_run_font(p1.add_run(value), size=9.5, color=INK)
    document.add_paragraph().paragraph_format.space_after = Pt(30)
    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(status.add_run("Локальная UAT-готовность: GO"), size=11, color=GREEN, bold=True)
    document.add_page_break()


def build_document() -> Document:
    document = Document()
    bullet_num_id, decimal_num_id = configure_document(document)
    props = document.core_properties
    props.title = "SBS AI ITSM — полное руководство"
    props.subject = "Пользовательское, административное и эксплуатационное руководство"
    props.author = "Smart Business Systems"
    props.keywords = "ITSM, Service Desk, CMDB, SLA, AI, руководство"

    add_cover(document)

    add_heading(document, "Содержание", 1)
    contents = [
        "1. Назначение платформы и границы документа",
        "2. Роли, права и зоны ответственности",
        "3. Быстрый старт и вход в систему",
        "4. Интерфейс, навигация и общие кнопки",
        "5. Инструкция для обычного пользователя",
        "6. Работа Service Desk и исполнителей",
        "7. Управление каталогом и сервисными заявками",
        "8. Изменения, проблемы, крупные инциденты и релизы",
        "9. Активы, CMDB и SLA",
        "10. База знаний и AI Copilot",
        "11. Аналитика, мониторинг и уведомления",
        "12. Автоматизация и интеграционные каналы",
        "13. Полное руководство администратора",
        "14. Системные настройки и SaaS Root",
        "15. Локальный запуск и ежедневная эксплуатация",
        "16. Перенос на сервер и production cutover",
        "17. Безопасность и устранение неисправностей",
        "18. Контрольные списки и словарь терминов",
        "Приложение A. Полная карта меню",
        "Приложение B. Статус готовности и проверок",
    ]
    add_bullets(document, contents, bullet_num_id)
    add_callout(document, "Как читать", "Обычному сотруднику достаточно разделов 3–5. Исполнителю — 6–12. Администратору — 13–17. Перед серверным запуском обязательно выполнить весь раздел 16.")

    add_heading(document, "1. Назначение платформы и границы документа", 1, new_page=True)
    add_heading(document, "1.1 Для чего нужна SBS AI ITSM", 2)
    add_paragraph(document, "SBS AI ITSM — единая multi-tenant платформа управления ИТ-услугами. Она объединяет обращение пользователей, работу Service Desk, учёт активов и связей, контроль SLA, управление изменениями и релизами, базу знаний, автоматизацию, интеграции и безопасное применение AI.")
    add_bullets(document, [
        ("Единая точка обращения. ", "Пользователь создаёт инцидент или выбирает услугу каталога и затем видит весь ход работы."),
        ("Управляемый процесс. ", "Назначения, согласования, сроки, комментарии и переходы статусов фиксируются в системе."),
        ("Прозрачность. ", "Менеджеры получают сводки, SLA, аналитику, календарь изменений и release evidence."),
        ("Контроль рисков. ", "RBAC, разделение обязанностей, tenant isolation, MFA, журнал аудита и защищённые секреты уменьшают вероятность ошибок и злоупотреблений."),
        ("Повторное использование знаний. ", "Статьи, KEDB и permission-aware AI помогают быстрее решать типовые обращения."),
        ("Интеграция. ", "Email, Microsoft Teams, мониторинг, SCIM/Entra, webhooks и внешние системы подключаются через управляемые каналы."),
    ], bullet_num_id)
    add_heading(document, "1.2 Что входит в платформу", 2)
    add_table(document, ["Область", "Для чего используется"], [
        ("Service Desk", "Инциденты, обращения, комментарии, назначения, статусы и массовые операции."),
        ("Service Catalog", "Стандартизированные услуги, формы, согласования, стоимость, сроки и fulfillment-задачи."),
        ("Change / Release", "RFC, CAB/ECAB, календарь окон, пакеты релиза, readiness, Go/No-Go и rollback."),
        ("Problem / KEDB", "Поиск первопричин, повторяющиеся инциденты, обходные решения и Known Errors."),
        ("Asset / CMDB", "Активы, CI-классы, связи, импорт, discovery, reconciliation, impact и качество данных."),
        ("Knowledge / AI", "База знаний, RAG-поиск, классификация, governance и защищённые AI-действия."),
        ("Operations", "SLA, события мониторинга, уведомления, email, Teams, аналитика и отчёты."),
        ("Administration", "Организации, пользователи, роли, права, SSO/SCIM, MFA, аудит и конфигурация."),
    ], [2500, 6860])
    add_heading(document, "1.3 Текущая граница готовности", 2)
    add_callout(document, "Локальный контур", "Система готова для демонстрации, настройки процессов, обучения и UAT. Основные роли и сквозные процессы проверены.", kind="success")
    add_callout(document, "Публичный production", "До публикации на сервере необходимо настроить домен/TLS, новые секреты, PostgreSQL/Redis, MFA и реальные внешние интеграции. Локальные demo-пароли нельзя использовать с реальными данными.", kind="warning")

    add_heading(document, "2. Роли, права и зоны ответственности", 1, new_page=True)
    add_heading(document, "2.1 Принцип доступа", 2)
    add_paragraph(document, "Доступ определяется не только названием роли, а эффективным набором permissions. Пользователь может иметь несколько ролей; итоговые права объединяются. Backend повторно проверяет permission, tenant_id и область видимости каждой записи, поэтому скрытие кнопки в интерфейсе не является единственной защитой.")
    add_table(document, ["Роль", "Основное назначение", "Типичная зона работы"], [
        ("Requester", "Заказчик/сотрудник", "Собственные инциденты и заявки, каталог, знания, уведомления, RAG."),
        ("IT Agent", "Исполнитель Service Desk", "Назначенные тикеты, fulfillment, события, активы, runbooks."),
        ("IT Manager", "Руководитель ИТ-процессов", "Все очереди организации, SLA, изменения, релизы, аналитика, согласования."),
        ("Organization Admin", "Администратор организации", "Пользователи, роли, настройки tenant, аудит, интеграционные параметры."),
        ("Security Officer", "Контроль безопасности", "Аудит, сессии, MFA, login events, identity и security-аналитика."),
        ("Knowledge Manager", "Владелец знаний", "Статьи, публикация, локализация, AI governance и качество знаний."),
        ("SaaS Root", "Владелец платформы", "Все tenant, системная диагностика, AI provider, production-конфигурация."),
    ], [1900, 2600, 4860], font_size=8.8)
    add_heading(document, "2.2 Разделение обязанностей", 2)
    add_bullets(document, [
        "Автор рискованного изменения не должен утверждать собственный RFC.",
        "Автор или владелец релиза не принимает Go/No-Go по своему релизу.",
        "Публикация каталога и знаний должна следовать review-процессу, если включён four-eyes control.",
        "Секреты AI, email, Teams, SCIM и webhooks доступны только через специализированные защищённые поля и никогда не должны копироваться в комментарии или тикеты.",
        "SaaS Root работает между организациями только осознанно; Organization Admin ограничен своим tenant.",
    ], bullet_num_id)
    add_heading(document, "2.3 Если раздел или кнопка отсутствуют", 2)
    add_steps(document, [
        "Проверьте, под какой учётной записью выполнен вход и какая организация выбрана.",
        "Откройте «Моя учётная запись» и проверьте назначенные роли/права.",
        "Обновите страницу после изменения ролей или выполните повторный вход.",
        "Администратор должен выдать минимально необходимое permission и соответствующую область видимости, например tickets.scope.assigned или requests.scope.all.",
        "Не пытайтесь обходить ограничение прямым URL или изменением базы данных; отказ 403 является ожидаемой защитой.",
    ], decimal_num_id)

    add_heading(document, "3. Быстрый старт и вход в систему", 1, new_page=True)
    add_heading(document, "3.1 Адреса локального контура", 2)
    add_table(document, ["Назначение", "Адрес"], [
        ("Пользовательский интерфейс", "http://localhost:5173"),
        ("Страница входа", "http://localhost:5173/login"),
        ("Backend API", "http://localhost:8000"),
        ("OpenAPI / Swagger", "http://localhost:8000/docs"),
        ("Health endpoint", "http://localhost:8000/api/v1/health"),
    ], [2800, 6560])
    add_heading(document, "3.2 Демонстрационные учётные записи", 2)
    add_callout(document, "Важно", "Эти записи предназначены только для DEMO_MODE=true. Перед серверным запуском отключите demo mode, деактивируйте записи или замените все пароли.", kind="danger")
    add_table(document, ["Роль", "Логин", "Пароль"], [
        ("Пользователь", "requester@sbs.local", "Sbs!2026"),
        ("ИТ-менеджер", "manager@sbs.local", "Sbs!2026"),
        ("Администратор", "admin@sbs.local", "Sbs!2026"),
        ("SaaS Root", "root@sbs.local", "Root!2026"),
    ], [2300, 4300, 2760])
    add_heading(document, "3.3 Вход", 2)
    add_steps(document, [
        "Откройте страницу входа.",
        "Выберите язык интерфейса: русский, казахский или английский.",
        "Введите корпоративный email и пароль.",
        "Нажмите «Войти». При включённом MFA введите код приложения-аутентификатора или recovery code.",
        "После входа проверьте роль и организацию в блоке контекста сессии.",
    ], decimal_num_id)
    add_heading(document, "3.4 Выход и защита сессии", 2)
    add_bullets(document, [
        "Для выхода используйте кнопку «Выйти» в верхней панели или боковом меню.",
        "Не оставляйте привилегированную сессию открытой на общем компьютере.",
        "При подозрении на компрометацию смените пароль, завершите другие сессии в профиле и сообщите Security Officer.",
        "После изменения ролей рекомендуется выйти и войти заново, чтобы получить актуальный session context.",
    ], bullet_num_id)

    add_heading(document, "4. Интерфейс, навигация и общие кнопки", 1, new_page=True)
    add_heading(document, "4.1 Основные области экрана", 2)
    add_table(document, ["Область", "Что показывает"], [
        ("Боковое меню", "Только те модули, которые разрешены эффективными permissions."),
        ("Верхняя панель", "Название раздела, глобальный поиск, язык, профиль и выход."),
        ("Контекст сессии", "Роль, организация, количество непрочитанных уведомлений и состояние сессии."),
        ("Сводные карточки", "Ключевые показатели текущего модуля; нули не должны скрывать ошибку загрузки."),
        ("Фильтры и вкладки", "Сужают очередь или переключают функциональную область без изменения прав."),
        ("Query failure panel", "Показывает, какие данные не загрузились, и предлагает повторить запрос."),
    ], [2400, 6960])
    add_heading(document, "4.2 Что делают типовые кнопки", 2)
    add_table(document, ["Кнопка/действие", "Назначение и правило"], [
        ("Создать", "Открывает форму нового объекта. Обязательные поля отмечены и проверяются backend."),
        ("Сохранить", "Сохраняет черновик или настройки. Для sensitive fields пустое значение обычно означает «не изменять»."),
        ("Отправить / Submit", "Переводит черновик на следующий управляемый этап; действие может стать необратимым."),
        ("Назначить себе", "Берёт объект в работу, если роль имеет self-assign и объект доступен по scope."),
        ("Назначить", "Выбирает исполнителя/группу; требует отдельного permission."),
        ("Согласовать / Отклонить", "Создаёт неизменяемое решение с причиной. Самосогласование может быть запрещено."),
        ("Опубликовать", "Делает утверждённую версию доступной пользователям; опубликованную версию не редактируют на месте."),
        ("Повторить", "Повторяет безопасную/idempotent операцию после анализа причины сбоя."),
        ("Проверить подключение", "Выполняет тест провайдера, но не обязательно сохраняет или активирует настройки."),
        ("Сохранить и активировать", "Применяет проверенную конфигурацию. Для внешних провайдеров требуется реальный секрет."),
        ("Отмена", "Закрывает форму без сохранения либо запускает управляемую отмену объекта — смотрите подтверждение."),
    ], [2500, 6860], font_size=8.7)
    add_heading(document, "4.3 Статусы и сообщения", 2)
    add_bullets(document, [
        ("Зелёный/Success. ", "Операция подтверждена системой или внешним провайдером."),
        ("Жёлтый/Warning. ", "Есть риск, просрочка, ожидание или неполная готовность."),
        ("Красный/Error. ", "Запрос не выполнен; прочитайте сообщение и correlation ID, не нажимайте повтор бесконечно."),
        ("SIMULATED/Mock. ", "Внешней доставки или LLM-вызова не было; результат предназначен для локальной проверки."),
        ("HTTP 409. ", "Состояние устарело или нарушен lifecycle. Обновите объект и повторно оцените действие."),
        ("HTTP 403. ", "Недостаточно прав или объект вне tenant/scope; это не техническая поломка."),
    ], bullet_num_id)

    add_heading(document, "5. Инструкция для обычного пользователя", 1, new_page=True)
    add_heading(document, "5.1 Когда выбирать каталог, а когда инцидент", 2)
    add_table(document, ["Ситуация", "Куда обращаться"], [
        ("Нужно предоставить стандартную услугу: VPN, ноутбук, ПО, общий mailbox", "Каталог услуг → Создать запрос"),
        ("Что-то сломалось или перестало работать", "Инциденты → Создать инцидент"),
        ("Нужно узнать инструкцию", "База знаний или AI Copilot"),
        ("Нужно проверить ход своей заявки", "Запросы услуг или Инциденты"),
        ("Нужно ответить исполнителю", "Откройте объект и добавьте комментарий/уточнение"),
    ], [4000, 5360])
    add_heading(document, "5.2 Подать заявку через каталог", 2)
    add_steps(document, [
        "Откройте «Каталог услуг».",
        "Найдите услугу поиском, фильтром категории или избранным.",
        "Откройте «Подробнее» и проверьте срок, владельца, стоимость, риск и необходимость согласования.",
        "Нажмите «Создать запрос».",
        "Заполните обязательные поля понятными деловыми формулировками; не помещайте пароли и API keys в обычные поля.",
        "При необходимости приложите разрешённые файлы и подтвердите отправку.",
        "Откройте «Запросы услуг» и проверьте номер REQ/RITM, статус, согласование и timeline.",
    ], decimal_num_id)
    add_heading(document, "5.3 Создать инцидент", 2)
    add_steps(document, [
        "Откройте «Инциденты» и нажмите кнопку создания.",
        "Укажите короткий заголовок: что не работает и где.",
        "В описании укажите время начала, затронутых пользователей, устройство/сервис и уже выполненные проверки.",
        "Выберите категорию и приоритет по фактическому влиянию. Не завышайте приоритет ради ускорения.",
        "Добавьте вложения без секретов и персональных данных, если они действительно помогают диагностике.",
        "Создайте инцидент и сохраните его номер SD-… для дальнейшей коммуникации.",
    ], decimal_num_id)
    add_heading(document, "5.4 Отслеживание и коммуникация", 2)
    add_bullets(document, [
        "Проверяйте статус и последние активности в самом объекте, а не только по email.",
        "Отвечайте на запрос уточнений через комментарий к тому же объекту, не создавая дубликат.",
        "Если проблема расширилась, добавьте факты и новых затронутых пользователей; исполнитель пересмотрит приоритет.",
        "После решения подтвердите результат. Если проблема осталась, используйте разрешённый возврат/reopen с объяснением.",
        "В «Уведомлениях» отмечайте прочитанные события и настраивайте предпочтения доставки.",
    ], bullet_num_id)
    add_heading(document, "5.5 База знаний и AI Copilot", 2)
    add_steps(document, [
        "Сначала выполните поиск в базе знаний по симптомам или названию сервиса.",
        "Оцените применимость статьи, дату публикации и область видимости.",
        "В Copilot формулируйте вопрос без паролей и секретов. RAG вернёт только источники, доступные вашей роли.",
        "Проверяйте ответ по ссылкам на источники. AI-рекомендация не отменяет утверждённые инструкции и change process.",
        "Оставьте feedback, если статья или ответ не помогли — это улучшает базу знаний.",
    ], decimal_num_id)

    add_heading(document, "6. Работа Service Desk и исполнителей", 1, new_page=True)
    add_heading(document, "6.1 Очередь инцидентов", 2)
    add_steps(document, [
        "Откройте «Инциденты» и выберите разрешённую очередь: назначенные, все доступные или requester scope.",
        "Отфильтруйте по статусу, приоритету, категории, SLA и исполнителю.",
        "Откройте карточку и проверьте requester, историю, затронутый актив/сервис, SLA и связанные объекты.",
        "Назначьте на подходящую группу/исполнителя или используйте «Назначить себе».",
        "Переведите в работу, фиксируйте диагностику и коммуникацию в activity stream.",
        "При ожидании укажите причину. Не используйте ожидание только для остановки SLA, если политика этого не допускает.",
        "Перед Resolve запишите решение и проверку результата. Close выполняйте согласно политике подтверждения.",
    ], decimal_num_id)
    add_heading(document, "6.2 Приоритет", 2)
    add_table(document, ["Приоритет", "Типовой смысл"], [
        ("CRITICAL", "Критичный сервис недоступен многим пользователям, серьёзный бизнес/безопасностный ущерб."),
        ("HIGH", "Сильное влияние, нет приемлемого обходного решения или затронута важная группа."),
        ("MEDIUM", "Ограниченное влияние, работа возможна частично или есть обходное решение."),
        ("LOW", "Незначительное влияние, информационный запрос или улучшение без срочности."),
    ], [1800, 7560])
    add_heading(document, "6.3 Связи и эскалация", 2)
    add_bullets(document, [
        "Свяжите повторяющиеся инциденты с Problem record, а не копируйте RCA в каждый тикет.",
        "Свяжите плановое исправление с Change/RFC.",
        "Если влияние соответствует major incident, используйте управляемое объявление крупного инцидента.",
        "При проблеме с активом привяжите CI; impact analysis покажет зависимости.",
        "При угрозе SLA уведомите владельца очереди заранее, не после breach.",
    ], bullet_num_id)
    add_heading(document, "6.4 Массовые действия", 2)
    add_paragraph(document, "Mass update выполняется через предварительный server-side preview и явное подтверждение. Перед подтверждением проверьте количество объектов, tenant, новое значение и возможные конфликты. Нельзя смешивать объекты разных организаций или выполнять массовое закрытие без обоснования.")

    add_heading(document, "7. Управление каталогом и сервисными заявками", 1, new_page=True)
    add_heading(document, "7.1 Структура каталога", 2)
    add_table(document, ["Уровень", "Назначение"], [
        ("Категория", "Навигационная группа, например «Рабочее место»."),
        ("Сервис", "Бизнес/ИТ-сервис и его владелец."),
        ("Offering", "Вариант предоставления сервиса."),
        ("Catalog Item", "То, что пользователь может заказать."),
        ("Form Version", "Версионированная схема полей, условий и валидации."),
        ("Entitlement", "Правила, кто видит и может заказать item."),
    ], [2300, 7060])
    add_heading(document, "7.2 Создание и публикация услуги", 2)
    add_steps(document, [
        "В режиме управления создайте/выберите категорию, сервис и offering.",
        "Создайте item-код, название, описание, владельца, fulfillment group, срок, стоимость, риск и approval policy.",
        "Создайте форму: ключи, типы, подписи на трёх языках, обязательность, условия, ограничения файлов и sensitive behavior.",
        "Настройте entitlement по ролям, группам, подразделениям и датам действия.",
        "Проверьте preview от имени разрешённого и запрещённого пользователя.",
        "Передайте Draft в In Review; reviewer не должен быть последним автором при four-eyes policy.",
        "Опубликуйте утверждённую версию и выполните один тестовый request.",
        "Для прекращения новых заказов используйте Retire, сохраняя историю существующих запросов.",
    ], decimal_num_id)
    add_heading(document, "7.3 Жизненный цикл заявки", 2)
    add_paragraph(document, "Нормальный поток: Submitted → Approval (если требуется) → Fulfillment → Completed. Допустимые управляемые исходы: Rejected, Rework, Cancelled и Failed.")
    add_steps(document, [
        "Проверьте версию item/form, entitlement, стоимость, SLA и requester.",
        "Согласующий принимает решение только в своём scope и указывает причину.",
        "Fulfillment manager назначает задачи подходящим пользователям/группам.",
        "Исполнитель запускает задачу, фиксирует результат и завершает только при наличии обязательных outputs.",
        "При rework укажите конкретное поле/задачу и сохраните прежнюю activity.",
        "Request автоматически/управляемо завершается только после терминального состояния всех обязательных задач.",
    ], decimal_num_id)

    add_heading(document, "8. Изменения, проблемы, крупные инциденты и релизы", 1, new_page=True)
    add_heading(document, "8.1 Change Management", 2)
    add_paragraph(document, "Изменение (RFC) используется для контролируемой модификации production-сервиса. Обязательны бизнес-обоснование, implementation, test, rollback и validation plans, оценка влияния и связанные CI.")
    add_steps(document, [
        "Создайте RFC в Draft и укажите тип Standard, Normal или Emergency.",
        "Свяжите затронутые сервисы, активы, инциденты и запросы.",
        "Заполните пять контрольных планов и отправьте на Assessment.",
        "Проверьте автоматически рассчитанный risk score и необходимость CAB/ECAB.",
        "Получите независимое решение Approve/Reject; автор не согласует собственный RFC.",
        "Назначьте окно. Пересечение по одному CI вернёт HTTP 409 — выберите другое время.",
        "Во время окна переведите в Implementing, выполните план и соберите evidence.",
        "При успехе выполните Review и Complete; при сбое — Failed/Rolled Back с причиной и результатом rollback.",
    ], decimal_num_id)
    add_heading(document, "8.2 Календарь изменений", 2)
    add_bullets(document, [
        "Используйте календарь для поиска пересечений, freeze periods и общей нагрузки на сервисы.",
        "Не удаляйте затронутый CI только ради обхода collision guard.",
        "Перед началом окна обновите RFC: version conflict означает, что объект изменил другой оператор.",
    ], bullet_num_id)
    add_heading(document, "8.3 Problem Management и KEDB", 2)
    add_steps(document, [
        "Создайте Problem для повторяющихся или значимых инцидентов и свяжите исходные записи.",
        "Проведите structured RCA: факты, временная линия, гипотезы, причина и подтверждение.",
        "Определите workaround и, если он полезен операторам, опубликуйте Known Error в KEDB.",
        "Создайте corrective/preventive actions, назначьте владельцев и сроки.",
        "Свяжите постоянное исправление с RFC и закройте Problem только после проверки результата.",
    ], decimal_num_id)
    add_heading(document, "8.4 Major Incident", 2)
    add_steps(document, [
        "Объявите major incident из подходящего критичного инцидента или создайте управляемую запись.",
        "Назначьте Incident Commander, technical lead и communication owner.",
        "Ведите единую timeline, решения, затронутые сервисы и коммуникации.",
        "Подключите Teams collaboration/war room только через настроенный production workflow.",
        "После восстановления зафиксируйте validation, link to Problem/RCA и итоговый обзор.",
    ], decimal_num_id)
    add_heading(document, "8.5 Release Management", 2)
    add_steps(document, [
        "Создайте release train с сервисом, semantic version, окном, scope, validation, rollback и communication plans.",
        "Добавьте только утверждённые RFC и все versioned artifacts с URI, SHA-256 и build reference.",
        "Независимый reviewer проверяет checksum, зависимости и manual gates.",
        "Переведите релиз в Ready только при пройденных обязательных readiness gates.",
        "Независимый decision maker выбирает GO, NO_GO или CONDITIONAL. Автор/владелец релиза не принимает решение.",
        "Продвигайте среды по порядку Test → Staging → Production; Production не может быть первым target.",
        "После deployment приложите logs, smoke/validation evidence и только затем публикуйте Released.",
        "При failure остановите продвижение, выполните rollback, восстановите previous_version и свяжите Incident/Problem.",
    ], decimal_num_id)

    add_heading(document, "9. Активы, CMDB и SLA", 1, new_page=True)
    add_heading(document, "9.1 Активы и Configuration Items", 2)
    add_bullets(document, [
        "Asset Inventory отвечает за владение, местоположение, состояние, проверку, перемещение, списание и историю оборудования/ПО.",
        "CMDB добавляет управляемые CI-классы, обязательные атрибуты, lifecycle status и отношения между CI.",
        "Связь инцидента/RFC/Problem с CI позволяет анализировать влияние и историю изменений.",
    ], bullet_num_id)
    add_heading(document, "9.2 Создание и обновление CI", 2)
    add_steps(document, [
        "Выберите правильный CI-класс; обязательные атрибуты зависят от класса (например hostname для server/device).",
        "Заполните уникальный asset tag/identity, владельца, местоположение, lifecycle и технические данные.",
        "Добавьте отношения только разрешённого типа и направления; запрещённые циклы будут отклонены.",
        "После изменения критичных данных выполните verification и проверьте history.",
        "Не редактируйте classified lifecycle обходным полем; используйте специализированные операции dispose/restore/retire.",
    ], decimal_num_id)
    add_heading(document, "9.3 Импорт, discovery и reconciliation", 2)
    add_steps(document, [
        "Создайте source/connector и проверьте credentials без вывода секретов.",
        "Для файла сначала выполните Preview: mapping, validation errors, duplicate candidates и class requirements.",
        "Исправьте ошибки, затем выполните Commit с idempotency safeguards.",
        "Discovery запускайте на разрешённой области; новые identities проходят matching/review.",
        "Reconciliation применяет приоритет источников, обнаруживает конфликты и сохраняет before/after history.",
        "Проверьте CMDB Quality: обязательные поля, stale CI, duplicate, lifecycle mismatch и orphan relationships.",
    ], decimal_num_id)
    add_heading(document, "9.4 Impact Analysis", 2)
    add_paragraph(document, "Перед change, problem или major incident откройте impact panel выбранного CI. Проверьте upstream/downstream зависимости, критичность сервиса, активные инциденты и текущие изменения. Анализ является подсказкой и не заменяет решение владельца сервиса.")
    add_heading(document, "9.5 SLA / OLA", 2)
    add_bullets(document, [
        "SLA policy определяет response/resolution targets по типу объекта, приоритету, календарю и организации.",
        "Warning означает приближение срока, Breached — подтверждённое нарушение target.",
        "Pause разрешён только для предусмотренных состояний ожидания; переходы и причина аудитируются.",
        "Приоритеты используются в едином верхнем регистре: LOW, MEDIUM, HIGH, CRITICAL.",
        "Менеджер регулярно анализирует breached и at-risk объекты, владельцев и причины повторных нарушений.",
    ], bullet_num_id)

    add_heading(document, "10. База знаний и AI Copilot", 1, new_page=True)
    add_heading(document, "10.1 Управление знаниями", 2)
    add_steps(document, [
        "Создайте статью в Draft с понятной проблемой, областью применимости, шагами и проверкой результата.",
        "Укажите категорию, audience/permissions и владельца актуальности.",
        "Передайте на review; reviewer проверяет безопасность, точность и отсутствие секретов.",
        "Опубликуйте утверждённую версию. Изменения опубликованного содержания выпускайте новой версией.",
        "Используйте feedback и usage analytics для улучшения; устаревшее содержание архивируйте.",
    ], decimal_num_id)
    add_heading(document, "10.2 Возможности AI", 2)
    add_table(document, ["Функция", "Назначение", "Ограничение"], [
        ("Classification", "Категория/приоритет/рекомендация для текста обращения.", "Оператор проверяет результат."),
        ("Permission-aware RAG", "Ответ по доступным статьям, тикетам и CI.", "Недоступные источники не раскрываются."),
        ("Governance", "Версии prompt/policy, evaluation, approve/deploy и audit.", "Разделение ролей и evidence."),
        ("Runtime Controls", "Лимиты, privacy, residency, budget и provider status.", "Fail-closed при нарушении policy."),
        ("Guarded Actions", "Предложение изменения объекта, approval, execute и rollback.", "Нет скрытого автономного действия."),
    ], [2200, 3600, 3560], font_size=8.5)
    add_heading(document, "10.3 Настройка OpenAI", 2)
    add_callout(document, "Важно", "Подписка ChatGPT и OpenAI API — разные продукты. Для платформы нужен API key проекта OpenAI; логин/пароль ChatGPT использовать нельзя.", kind="warning")
    add_steps(document, [
        "Войдите как Organization Admin или SaaS Root и откройте «Администрирование → Настройки» либо системную AI configuration panel.",
        "Выберите provider OpenAI.",
        "Введите API key в защищённое поле, модель проекта и base URL https://api.openai.com/v1 (или утверждённый gateway).",
        "Оставьте PII redaction включённым.",
        "Нажмите «Проверить подключение» и прочитайте requested/effective provider, HTTP status и reason.",
        "После успешного теста нажмите «Сохранить и активировать».",
        "Проверьте AI Provider Status и выполните безопасный тест Copilot без персональных данных.",
    ], decimal_num_id)
    add_heading(document, "10.4 Настройка Gemini", 2)
    add_steps(document, [
        "Создайте ограниченный key в Google AI Studio/проекте Google.",
        "Выберите Gemini, укажите key и доступную проекту модель.",
        "Ограничьте key Generative Language API и backend-средой, где это возможно.",
        "Выполните connection test, затем сохраните и активируйте.",
        "Не помещайте key во frontend, localStorage, screenshot, тикет или этот документ.",
    ], decimal_num_id)
    add_heading(document, "10.5 Что означает ai_copilot_enabled=true", 2)
    add_paragraph(document, "Это только feature flag, разрешающий функции Copilot. Он не выбирает provider и не задаёт API key. Если real provider не настроен, система честно остаётся в Mock/Local Simulation. Mock поддерживает демонстрацию правил, но не является внешней LLM.")

    add_heading(document, "11. Аналитика, мониторинг и уведомления", 1, new_page=True)
    add_heading(document, "11.1 Dashboard и Analytics", 2)
    add_paragraph(document, "Dashboard показывает только доступные пользователю данные. В Analytics используются вкладки Executive, Tickets, SLA, Assets, Knowledge, AI, Security, Automation и Reports.")
    add_bullets(document, [
        "Executive: общий объём, backlog, динамика, critical risks и ключевые тренды.",
        "Tickets: создание/закрытие, категории, приоритеты, возраст и производительность очередей.",
        "SLA: warning, breach, соблюдение целей и причины нарушений.",
        "Assets: состояние, качество, lifecycle и распределение активов.",
        "Knowledge/AI: использование статей, feedback, provider execution и качество рекомендаций.",
        "Security: login/audit/session события в разрешённой области.",
        "Automation: runs, failures, approvals и эффективность workflow.",
        "Reports: сохранённые отчёты, snapshots и экспорт согласно permissions.",
    ], bullet_num_id)
    add_heading(document, "11.2 Monitoring и Event Operations", 2)
    add_steps(document, [
        "Создайте monitoring source с HMAC, bearer или provider-specific authentication.",
        "Отправьте тестовое событие и проверьте receipt, signature, parsing и tenant mapping.",
        "Контролируйте receipt queue, retries и dead letter; не переотправляйте неподтверждённые payload вслепую.",
        "В Event Operations нормализуйте, дедуплицируйте и коррелируйте события.",
        "Создавайте Incident только по политике; сохраняйте связь с source event.",
    ], decimal_num_id)
    add_heading(document, "11.3 Уведомления", 2)
    add_bullets(document, [
        "In-app уведомление — источник состояния даже при недоступном email/Teams.",
        "Пользователь управляет разрешёнными preferences; администратор — templates и каналы.",
        "Email Log показывает provider, статус и ошибку. Retry разрешён только для FAILED/BOUNCED и после анализа причины.",
        "Статус SIMULATED означает, что реальная отправка не выполнялась.",
    ], bullet_num_id)

    add_heading(document, "12. Автоматизация и интеграционные каналы", 1, new_page=True)
    add_heading(document, "12.1 Automation", 2)
    add_paragraph(document, "Раздел содержит Production Workflows, Overview, Rules, Executions, Runbooks, Approvals и Templates/Examples.")
    add_steps(document, [
        "Создайте rule/workflow в Draft и выберите точный trigger.",
        "Добавьте conditions и только разрешённые safe actions. Secret-like fields валидатор отклоняет.",
        "Настройте retry: ограниченное число попыток и backoff; workflow graph не должен содержать cycle.",
        "Проведите simulation без side effects и запросите review/publish.",
        "Следите за execution events, waits, approvals, retries и dead letter.",
        "Повторяйте/rollback только после проверки idempotency и текущего состояния объекта.",
    ], decimal_num_id)
    add_heading(document, "12.2 Интеграционная платформа", 2)
    add_steps(document, [
        "Создайте service account с минимальными permissions и tenant scope.",
        "Выпустите token, сохраните его во внешнем secret store и зафиксируйте дату rotation.",
        "Для outbound webhook задайте HTTPS URL, event allowlist, signing secret, retry и rate limit.",
        "Проверьте подпись и idempotency на стороне получателя.",
        "Наблюдайте deliveries, latency, retry и dead-letter; replay выполняйте только после устранения причины.",
        "При утечке немедленно revoke/rotate token или webhook secret и просмотрите audit trail.",
    ], decimal_num_id)
    add_heading(document, "12.3 Production Email Channel", 2)
    add_bullets(document, [
        "Реальный канал использует Microsoft Graph/Entra application и backend secrets.",
        "После создания выполните connection test, активируйте inbound/outbound и настройте Graph webhook/subscription.",
        "Inbound письмо связывается с существующим обращением или создаёт новое по политике; attachments проходят ограничения и безопасное хранение.",
        "Mock provider только симулирует; у SIMULATED нет provider message ID и delivery timestamp.",
    ], bullet_num_id)
    add_heading(document, "12.4 Microsoft Teams", 2)
    add_bullets(document, [
        "Создайте Teams Workflow/connector и храните webhook URL как secret.",
        "Настройте event types, severity, destination и message template.",
        "Для major incident используйте отдельный управляемый канал/war room и аудит действий.",
        "FAILED/DEAD_LETTER можно повторить после устранения причины; SIMULATED не означает HTTP-доставку.",
    ], bullet_num_id)

    add_heading(document, "13. Полное руководство администратора", 1, new_page=True)
    add_heading(document, "13.1 Вкладки администрирования", 2)
    add_table(document, ["Вкладка", "Назначение"], [
        ("Overview", "Сводка tenant, пользователей, ролей, безопасности и готовности."),
        ("Пользователи", "Создание, изменение профиля, несколько ролей, активация/деактивация."),
        ("Роли и права", "Permission catalog, поиск, группировка, dependency warnings и замена набора прав."),
        ("Аудит", "Tenant-scoped tamper-evident события, фильтры и проверка цепочки."),
        ("Настройки", "Typed settings, feature flags и отдельная AI provider panel."),
        ("Identity & SSO", "OIDC, identity provider, SCIM/Entra provisioning и linking."),
        ("Security", "Сессии, login events, MFA posture и привилегированные действия."),
        ("Организация", "Профиль/branding/локализация tenant; создание tenant доступно root."),
    ], [2200, 7160], font_size=8.8)
    add_heading(document, "13.2 Создать пользователя и назначить роли", 2)
    add_steps(document, [
        "Откройте «Администрирование → Пользователи».",
        "Нажмите создание пользователя и заполните email, полное имя, подразделение/должность при необходимости.",
        "Выберите одну или несколько ролей. Система объединит их permissions.",
        "Назначьте временный пароль согласно password policy и включите обязательную смену при первом входе, если доступно.",
        "Сохраните и проверьте пользователя в списке.",
        "Выполните тестовый вход или попросите владельца подтвердить доступ к необходимым разделам.",
        "При увольнении используйте deprovision/offboarding: передайте ownership, отзовите sessions и деактивируйте запись.",
    ], decimal_num_id)
    add_heading(document, "13.3 Создать или изменить роль", 2)
    add_steps(document, [
        "Скопируйте подходящую модель доступа или создайте tenant role с понятным кодом/описанием.",
        "Через поиск permissions выберите только необходимые capabilities.",
        "Для tickets.read обязательно задайте visibility scope: all, assigned или requester; для requests.read — all или requester.",
        "Проверьте предупреждения dependency guard и устраните их до сохранения.",
        "Назначьте роль тестовому пользователю, выполните позитивный и негативный сценарий.",
        "Не изменяйте системную роль ради одного пользователя, если безопаснее создать отдельную роль.",
    ], decimal_num_id)
    add_heading(document, "13.4 Организации", 2)
    add_paragraph(document, "SaaS Root создаёт организацию, задаёт slug/status и назначает первого Organization Admin. Дальнейшие пользователи, роли, настройки, каталоги и записи должны оставаться tenant-scoped. Перед операциями root всегда проверяет выбранную организацию.")
    add_heading(document, "13.5 Audit Trail", 2)
    add_steps(document, [
        "Выберите период, actor, action/module, tenant и correlation ID.",
        "Откройте событие и сопоставьте его с business history объекта.",
        "При расследовании экспортируйте только разрешённые данные и сохраните chain verification result.",
        "Никогда не исправляйте audit записи вручную; выявленное нарушение цепочки является security incident.",
    ], decimal_num_id)
    add_heading(document, "13.6 Identity Provisioning", 2)
    add_steps(document, [
        "Создайте SCIM/Entra connector, default role, fallback owner, IP allowlist и retry policy.",
        "Скопируйте выданный token один раз в Entra secret configuration; позже он не отображается полностью.",
        "Настройте group-to-role mapping и проверьте joiner/mover на тестовом пользователе.",
        "Для leaver выберите нового владельца; система передаст активные объекты и отзовёт сессии.",
        "Контролируйте events: APPLIED, RETRY_SCHEDULED, FAILED, DEAD_LETTER; ручной retry — после исправления причины.",
    ], decimal_num_id)
    add_heading(document, "13.7 Custom Fields и Configuration Packages", 2)
    add_bullets(document, [
        "Custom Field: создайте definition, data type, entity scope, validation, sensitivity, search/report flags; затем publish version.",
        "Sensitive field читается только с отдельным permission и не должен попадать в обычный audit/log/export.",
        "Configuration Package собирает управляемые настройки tenant, seal/export/import, deployment plan, approve/apply и rollback.",
        "Перед применением package выполните diff/plan и независимое approval; rollback использует сохранённое evidence.",
    ], bullet_num_id)

    add_heading(document, "14. Системные настройки и SaaS Root", 1, new_page=True)
    add_heading(document, "14.1 System Diagnostics", 2)
    add_table(document, ["Проверка", "Что означает"], [
        ("Health", "Базовый ответ приложения и версия."),
        ("Liveness", "Процесс жив и может отвечать."),
        ("Readiness", "Критичные зависимости готовы принимать трафик."),
        ("Deep health", "Расширенная диагностика базы, Redis, worker и компонентов."),
        ("Jobs summary/runtime", "Состояние очередей, исполнителей, retries и dead letter."),
        ("Outbox diagnostics", "Транзакционные события ожидают публикации или застряли."),
        ("AI provider status", "Configured/effective provider, readiness, simulation и reason без секрета."),
    ], [2500, 6860])
    add_heading(document, "14.2 Unified Configuration Center", 2)
    add_bullets(document, [
        "Показывает readiness domains и typed settings, а не только строки true/false.",
        "Изменение проводится с reason, actor, history и возможностью governed rollback.",
        "Sensitive values не возвращаются браузеру; UI показывает только configured boolean/hint.",
        "Feature flag и provider configuration — разные сущности. Например ai_copilot_enabled не хранит OpenAI key.",
    ], bullet_num_id)
    add_heading(document, "14.3 Безопасная работа с секретами", 2)
    add_bullets(document, [
        "В production предпочитайте Docker secrets, Vault/KMS или environment injection.",
        "Никогда не вставляйте ключи в чат, screenshot, браузерную консоль, тикет, wiki или Git.",
        "После ввода secret read API возвращает только признак configured и, где допустимо, короткий hint.",
        "При замене введите полный новый key; старое значение нельзя восстановить через UI.",
        "При подозрении на утечку сначала revoke/rotate у провайдера, затем обновите платформу и просмотрите audit.",
    ], bullet_num_id)

    add_heading(document, "15. Локальный запуск и ежедневная эксплуатация", 1, new_page=True)
    add_heading(document, "15.1 Быстрый запуск Windows", 2)
    add_code_block(document, [r"cd C:\projects\sbs-ai-itsm-foundation-001", r"wscript.exe scripts\start-local-demo.vbs"])
    add_paragraph(document, "Скрипт скрыто запускает backend на 127.0.0.1:8000 и frontend на 127.0.0.1:5173, если порты ещё не заняты. Локальные данные хранятся в постоянной SQLite-базе:")
    add_code_block(document, [r"%LOCALAPPDATA%\Temp\sbs-ai-itsm-visible.db"])
    add_heading(document, "15.2 Отдельный запуск компонентов", 2)
    add_code_block(document, [r"scripts\start-local-backend.cmd", r"scripts\start-local-frontend.cmd"])
    add_heading(document, "15.3 Проверка после запуска", 2)
    add_steps(document, [
        "Откройте health endpoint и убедитесь, что status=ok.",
        "Откройте /login и выполните вход requester.",
        "Проверьте каталог, создание тестового инцидента и отображение уведомлений.",
        "Войдите manager и проверьте очереди, self-assignment и SLA.",
        "Войдите admin/root и проверьте System readiness, jobs/outbox и provider status.",
    ], decimal_num_id)
    add_heading(document, "15.4 Диагностика портов", 2)
    add_code_block(document, [r"Get-NetTCPConnection -LocalPort 8000,5173 -State Listen", r"Get-Process -Id <PID>"])
    add_callout(document, "Остановка", "Останавливайте только PID, который после проверки принадлежит локальному python/uvicorn или node/vite этого проекта. Не используйте массовое завершение всех Python/Node процессов.", kind="warning")
    add_heading(document, "15.5 Логи", 2)
    add_table(document, ["Файл", "Назначение"], [
        (".local-backend.out.log", "Структурированные HTTP и application события."),
        (".local-backend.err.log", "Старт Uvicorn, traceback и stderr."),
        (".local-frontend.out.log", "Vite output."),
        (".local-frontend.err.log", "Proxy errors и browser console warnings/errors."),
    ], [3600, 5760])

    add_heading(document, "16. Перенос на сервер и production cutover", 1, new_page=True)
    add_callout(document, "Решение", "Код и compose-конфигурация проходят локальный release gate, но серверный запуск допускается только после выполнения всех шагов этой главы.", kind="danger")
    add_heading(document, "16.1 Обязательные входные данные", 2)
    add_bullets(document, [
        "Домен, DNS, TLS certificate и reverse proxy/load balancer.",
        "Production-серверы/VM, firewall rules, outbound HTTPS, storage и backup policy.",
        "Новые JWT/database/Redis/webhook/OIDC/AI secrets, созданные специально для production.",
        "PostgreSQL и Redis с persistent volumes и мониторингом.",
        "Решение по корпоративному OIDC/Entra, MFA и привилегированным аккаунтам.",
        "Реальные OpenAI/Gemini, Microsoft Graph, Teams и monitoring credentials — только для нужных интеграций.",
    ], bullet_num_id)
    add_heading(document, "16.2 Production-параметры", 2)
    add_table(document, ["Параметр", "Требование"], [
        ("DEMO_MODE", "false"),
        ("SEED_DEMO_CATALOG", "false, если demo-услуги не нужны"),
        ("RUN_STARTUP_DDL", "false; схему меняет только Alembic"),
        ("JWT_SECRET_KEY", "Новый сильный уникальный secret"),
        ("DATABASE_URL", "Production PostgreSQL, не SQLite"),
        ("REDIS_URL", "Production Redis"),
        ("CORS / frontend origin", "Только реальный HTTPS origin"),
        ("Trusted hosts", "Только утверждённые домены/proxy"),
        ("JOBS_EXECUTOR_MODE", "redis с отдельным worker"),
    ], [3000, 6360])
    add_heading(document, "16.3 Рекомендуемая последовательность", 2)
    add_steps(document, [
        "Создайте .env.production из шаблона и заполните его через защищённый канал.",
        "Запустите scripts/check-production-env.sh и устраните все FAIL, не выводя secret values.",
        "Проверьте docker-compose.prod.yml и docker-compose.bootstrap.yml командой config.",
        "Поднимите PostgreSQL/Redis и сделайте отдельную backup/checkpoint точку перед миграцией.",
        "Выполните alembic upgrade head и подтвердите единственную голову 20260814_0073 для текущей версии.",
        "Запустите API, worker, scheduler и frontend через production Compose.",
        "Проверьте health/liveness/readiness/deep health, queues/outbox и отсутствие 5xx.",
        "Создайте реальные организации и администраторов, включите MFA, смените/удалите demo accounts.",
        "Настройте и протестируйте только необходимые интеграции.",
        "Выполните browser/API smoke, tenant-isolation acceptance, нагрузочный тест и security scan.",
        "Получите формальное Go/No-Go и только затем допускайте реальных пользователей.",
    ], decimal_num_id)
    add_heading(document, "16.4 Команды deployment", 2)
    add_code_block(document, [
        "cp .env.production.example .env.production",
        "bash scripts/check-production-env.sh",
        "docker compose -f docker-compose.prod.yml --env-file .env.production config",
        "docker compose -f docker-compose.prod.yml --env-file .env.production up -d --build",
        "docker compose -f docker-compose.prod.yml --env-file .env.production exec backend alembic upgrade head",
    ])
    add_heading(document, "16.5 Acceptance после переноса", 2)
    add_bullets(document, [
        "Frontend работает только по HTTPS; HTTP корректно перенаправляется.",
        "Readiness подтверждает Postgres, Redis, worker и критичные зависимости.",
        "Requester не видит чужие tenant/records; admin ограничен своей организацией.",
        "MFA обязателен для admin/root/security; recovery процедура проверена.",
        "Email/Teams/AI помечают реальную доставку только при подтверждении провайдера.",
        "Созданы incident, service request, change и release smoke records с audit evidence.",
        "Backup и restore выполнены на тестовом наборе данных; мониторинг и alerts поступают ответственным.",
    ], bullet_num_id)

    add_heading(document, "17. Безопасность и устранение неисправностей", 1, new_page=True)
    add_heading(document, "17.1 Минимальные правила безопасности", 2)
    add_bullets(document, [
        "Least privilege: выдавайте право только для реальной рабочей обязанности.",
        "Привилегированные пользователи используют MFA и отдельные персональные записи, а не общий admin login.",
        "Секреты хранятся вне Git и frontend; их значения не входят в audit trail.",
        "Нельзя отключать tenant scope, four-eyes, evidence или lifecycle ради ускорения операции.",
        "Перед retry/rollback проверяйте idempotency, version и текущее состояние объекта.",
        "Security incident включает сохранение evidence, revoke/rotate credentials, анализ audit и коммуникацию владельцам.",
    ], bullet_num_id)
    add_heading(document, "17.2 Частые проблемы", 2)
    add_table(document, ["Симптом", "Что проверить"], [
        ("Страница входа не открывается", "Порт 5173, node/vite process, .local-frontend.err.log."),
        ("API недоступен", "Порт 8000, health, backend log, SQLite/Postgres URL и миграции."),
        ("Везде красный блок «данные недоступны»", "Backend health, proxy /api, авторизация, конкретные failed queries и Retry."),
        ("403", "Роль, permission, tenant и record scope; повторный вход после изменения ролей."),
        ("Каталог пуст", "Published lifecycle, parent service/offering, entitlement, effective dates и SEED_DEMO_CATALOG."),
        ("Заявка не отправляется", "Form version/hash, required/conditional fields, attachments, entitlement и idempotency."),
        ("409 при изменении", "Обновите объект: stale version, lifecycle conflict, overlapping window или active deployment."),
        ("AI показывает true, но не работает", "true — feature flag. Откройте AI provider panel, введите key/model, Test и Activate."),
        ("AI 401/403/404/429", "Недействительный key/доступ/model/quota; проверьте provider project и ограничения."),
        ("Email/Teams = SIMULATED", "Работает mock. Настройте реальный Graph/Workflow connector; не трактуйте как доставку."),
        ("Jobs остаются queued", "Redis, worker, outbox age, queue name, scheduler и executor mode."),
        ("SCIM event FAILED", "Connector status, token/IP allowlist, mapping, default role/fallback owner и retry event."),
        ("Миграция не стартует", "alembic heads/current, DATABASE_URL, права DB и последовательность bootstrap."),
    ], [3100, 6260], font_size=8.2)
    add_heading(document, "17.3 Что приложить к обращению о сбое", 2)
    add_bullets(document, [
        "Время и часовой пояс, URL раздела, роль и организация без пароля.",
        "Номер объекта (SD/REQ/RITM/RFC/Release) и точное действие.",
        "Текст ошибки, HTTP status и correlation ID.",
        "Безопасный screenshot без секретов и персональных данных.",
        "Последние релевантные строки лога; не прикладывайте весь .env или provider key.",
    ], bullet_num_id)

    add_heading(document, "18. Контрольные списки и словарь терминов", 1, new_page=True)
    add_heading(document, "18.1 Ежедневный checklist Service Desk", 2)
    add_bullets(document, [
        "Проверить critical/high backlog и неназначенные обращения.",
        "Проверить SLA warning/breached и объекты в ожидании без причины.",
        "Проверить major incidents, monitoring events и failed deliveries.",
        "Проверить fulfillment tasks, approvals и rework.",
        "Обновить комментарии/статусы и связать повторяющиеся инциденты с Problem.",
    ], bullet_num_id)
    add_heading(document, "18.2 Еженедельный checklist менеджера", 2)
    add_bullets(document, [
        "Backlog age, SLA trend, повторные категории, major incidents и Problem actions.",
        "Change calendar, конфликтующие окна, CAB/ECAB и upcoming releases.",
        "Release success/failure/rollback, readiness drift и невыполненные gates.",
        "CMDB quality, stale assets, unverified CI и reconciliation conflicts.",
        "Automation failures, dead letter, integration health и notification delivery.",
        "Knowledge gaps, Copilot feedback и provider status/cost controls.",
    ], bullet_num_id)
    add_heading(document, "18.3 Ежемесячный checklist администратора", 2)
    add_bullets(document, [
        "Проверить активных пользователей, лишние роли и неиспользуемые privileged accounts.",
        "Проверить MFA posture, login events, sessions и audit chain integrity.",
        "Rotate истекающие tokens/secrets и проверить webhook/Graph/SCIM subscriptions.",
        "Проверить readiness, jobs/outbox, storage, backup/restore evidence и certificate expiry.",
        "Проверить catalog/knowledge ownership, локализацию и устаревшие версии.",
    ], bullet_num_id)
    add_heading(document, "18.4 Словарь", 2)
    add_table(document, ["Термин", "Определение"], [
        ("ITSM", "Управление ИТ-услугами как набор процессов, ролей и измеримых результатов."),
        ("Incident", "Нарушение или снижение качества услуги."),
        ("Service Request", "Стандартный запрос на услугу/доступ/оборудование."),
        ("Problem", "Причина одного или нескольких инцидентов."),
        ("Known Error / KEDB", "Известная проблема с документированным workaround."),
        ("RFC / Change", "Управляемое изменение сервиса или инфраструктуры."),
        ("CAB / ECAB", "Орган согласования обычных/аварийных изменений."),
        ("Release", "Управляемый набор изменений и versioned artifacts для продвижения по средам."),
        ("CI", "Configuration Item — управляемая единица в CMDB."),
        ("CMDB", "База конфигурационных единиц, атрибутов и отношений."),
        ("SLA / OLA", "Цель услуги для заказчика / внутренняя операционная цель."),
        ("RAG", "Генерация ответа с поиском по разрешённым внутренним источникам."),
        ("RBAC", "Управление доступом через роли и permissions."),
        ("Tenant", "Изолированная организация внутри SaaS-платформы."),
        ("Idempotency", "Повтор одинаковой команды не создаёт повторный бизнес-эффект."),
        ("Outbox", "Транзакционная очередь событий, ожидающих публикации."),
        ("Dead letter", "Очередь операций, исчерпавших допустимые retries."),
        ("Mock / SIMULATED", "Локальная имитация без подтверждённого внешнего действия."),
    ], [2400, 6960], font_size=8.5)

    add_heading(document, "Приложение A. Полная карта меню", 1, new_page=True)
    menu_rows = [
        ("/account", "Моя учётная запись", "Профиль, пароль, MFA и сессии."),
        ("/dashboard", "Обзор", "Ролевой operational dashboard."),
        ("/catalog", "Каталог услуг", "Поиск/заказ услуг и управление каталогом."),
        ("/requests", "Запросы услуг", "Заявки, approval, fulfillment и timeline."),
        ("/tickets", "Инциденты", "Очереди Service Desk и карточка тикета."),
        ("/major-incidents", "Крупные инциденты", "Major incident command, timeline и коммуникации."),
        ("/events", "События", "Monitoring receipts, events и создание incident."),
        ("/changes", "Изменения", "RFC lifecycle и history."),
        ("/change-calendar", "Календарь изменений", "Окна, пересечения и планирование."),
        ("/releases", "Релизы", "Release trains, gates, Go/No-Go и deployments."),
        ("/problems", "Проблемы и KEDB", "Problem lifecycle, workaround и Known Errors."),
        ("/problem-governance", "RCA и тренды", "Recurring analytics, RCA и corrective actions."),
        ("/assets", "Активы", "Inventory, CMDB, import, discovery и quality."),
        ("/sla", "SLA", "Политики, цели, предупреждения и breaches."),
        ("/knowledge", "База знаний", "Статьи, review, publish и feedback."),
        ("/copilot", "AI Copilot", "AI analysis, RAG, governance и actions."),
        ("/notifications", "Уведомления", "Inbox, preferences и templates."),
        ("/notifications/email-log", "Email Log", "Статусы outbound email и retry."),
        ("/analytics", "Аналитика", "Executive и доменные dashboards/reports."),
        ("/monitoring", "Monitoring", "System-level monitoring view."),
        ("/automation", "Автоматизация", "Rules, workflows, runbooks, executions и approvals."),
        ("/integrations", "Интеграции", "Legacy/demo и production integration platform."),
        ("/admin", "Администрирование", "Пользователи, роли, audit, settings, identity, security, tenant."),
        ("/admin/system", "System", "Health, jobs/outbox и AI provider diagnostics."),
        ("/identity-provisioning", "Identity Provisioning", "SCIM/Entra joiner-mover-leaver."),
        ("/email-operations", "Почтовый канал", "Graph channels, sync, webhook и delivery."),
        ("/teams-collaboration", "Microsoft Teams", "Connectors/workflows и deliveries."),
        ("/admin/custom-fields", "Настраиваемые поля", "Definitions, versions и values."),
        ("/admin/configuration-packages", "Пакеты конфигурации", "Build, seal, export/import, deploy и rollback."),
    ]
    add_table(document, ["Маршрут", "Раздел", "Назначение"], menu_rows, [2200, 2700, 4460], font_size=7.8)

    add_heading(document, "Приложение B. Статус готовности и проверок", 1, new_page=True)
    add_callout(document, "Итог", "Локальная демонстрация и UAT — GO. Публичный production — после server cutover из раздела 16.", kind="success")
    add_table(document, ["Контроль", "Результат"], [
        ("Backend", "83 тестовых модуля; 729 собранных сценариев."),
        ("Release gate", "27 из 27 проверок PASS."),
        ("API security", "709 endpoints; 702 protected; 7 governed public."),
        ("OpenAPI", "591 paths."),
        ("Alembic", "Единая голова 20260814_0073."),
        ("Frontend", "TypeScript PASS; Vite production build; 140 modules."),
        ("Accessibility", "75 файлов; 16 dialogs; нет unnamed/no-focus dialogs."),
        ("Controls", "651 buttons и 39 links проверены."),
        ("Live smoke", "Frontend 200; backend health ok; 5xx/traceback/console/proxy errors — 0."),
        ("Языки", "Русский, казахский и английский в общем UI и проверенных основных потоках."),
    ], [2800, 6560])
    add_paragraph(document, "Машиночитаемое evidence: docs/reports/PRODUCTION-READINESS-AUDIT-2026-08-13.json", italic=True, color=MUTED)
    add_paragraph(document, "Evidence SHA-256: 17a683fba3138f7f6083183f97fe6f76d49aa6a2739b109be823bbefa49d2107", italic=True, color=MUTED)
    add_callout(document, "Примечание о локализации", "Основные пользовательские потоки и общий интерфейс проверены на трёх языках. Перед публичным запуском рекомендуется лингвистическая приёмка редких административных экранов носителями русского и казахского языков.", kind="info")

    end = document.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end.paragraph_format.space_before = Pt(24)
    set_run_font(end.add_run("Конец руководства"), size=10, color=MUTED, italic=True)
    return document


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
