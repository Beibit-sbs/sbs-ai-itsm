from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


DOCUMENT = Path(
    r"C:\projects\sbs-ai-itsm-foundation-001\docs\manuals\SBS-AI-ITSM-FULL-GUIDE-RU-v0.1.0.docx"
)
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def main() -> None:
    document = Document(DOCUMENT)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )

    with zipfile.ZipFile(DOCUMENT) as archive:
        zip_error = archive.testzip()
        root = etree.fromstring(archive.read("word/document.xml"))
        office_xml = b"\n".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )

    tables = root.xpath(".//w:tbl", namespaces=NS)
    widths = [
        table.xpath("./w:tblPr/w:tblW/@w:w", namespaces=NS) for table in tables
    ]
    indents = [
        table.xpath("./w:tblPr/w:tblInd/@w:w", namespaces=NS) for table in tables
    ]
    grids = [
        sum(
            int(width)
            for width in table.xpath(
                "./w:tblGrid/w:gridCol/@w:w", namespaces=NS
            )
        )
        for table in tables
    ]

    result = {
        "file_bytes": DOCUMENT.stat().st_size,
        "zip_error": zip_error,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "headings": sum(
            1
            for paragraph in document.paragraphs
            if paragraph.style and paragraph.style.name.startswith("Heading")
        ),
        "numbered_paragraphs": len(root.xpath(".//w:numPr", namespaces=NS)),
        "repeating_table_headers": len(
            root.xpath(".//w:tblHeader", namespaces=NS)
        ),
        "page_fields": office_xml.count(b"PAGE"),
        "fixed_row_heights": len(root.xpath(".//w:trHeight", namespaces=NS)),
        "all_table_widths_9360": all(width == ["9360"] for width in widths),
        "all_table_indents_120": all(indent == ["120"] for indent in indents),
        "all_grids_9360": all(grid == 9360 for grid in grids),
        "placeholders": re.findall(
            r"\[\[.*?\]\]|\b(?:TODO|TBD)\b", text, flags=re.IGNORECASE
        ),
        "mojibake_markers": [
            marker
            for marker in ("Рђ", "Рџ", "СЃ", "Рё", "Р°")
            if marker in text
        ],
        "required_sections": all(
            section in text
            for section in (
                "Календарь изменений",
                "AI Copilot",
                "OpenAI",
                "Gemini",
                "Приложение A",
                "Приложение B",
            )
        ),
    }
    print(result)


if __name__ == "__main__":
    main()
